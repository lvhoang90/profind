#!/usr/bin/env python3
"""Nhập danh mục nguồn Scopus (tệp Excel "ext_list_<tháng>_<năm>.xlsx" của Elsevier, sheet "Scopus Sources ...") thành data/scopus-sources.json.
Dùng:  python3 -I scripts/import-scopus-list.py <đường dẫn xlsx> [--out data/scopus-sources.json]
Chỉ giữ ISSN của công trình đã tính điểm theo Scopus hoặc có hạng hiển thị (scoreKind = scopus hoặc jq; từ public/data/works/*.json) để tệp gọn.
Mỗi ISSN (bỏ gạch nối, 8 ký tự) -> danh sách [tên nguồn, Active|Inactive, Coverage (chuỗi các năm/khoảng năm), loại nguồn]; một ISSN có thể thuộc nhiều bản ghi (tên cũ/mới)."""
import sys, json, glob, os
import openpyxl
src = sys.argv[1]; out = sys.argv[sys.argv.index('--out') + 1] if '--out' in sys.argv else 'data/scopus-sources.json'
need = set()
for f in glob.glob('public/data/works/w*.json'):
    for ws in json.load(open(f, encoding='utf8')).values():
        for w in ws:
            i = (w.get('issn') or '').replace('-', '').upper()
            if i and (w.get('scoreKind') == 'scopus' or w.get('jq')): need.add(i.zfill(8))
wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
sheet = next(n for n in wb.sheetnames if n.lower().startswith('scopus sources'))
rows = wb[sheet].iter_rows(values_only=True); next(rows)
res = {}; built = sheet
for r in rows:
    title, issn, eissn, act, cov, typ = r[1], r[2], r[3], r[4], r[5], r[12]
    rec = [str(title or ''), str(act or ''), str(cov or '').replace(' ', ''), str(typ or '')]
    for i in (issn, eissn):
        i = str(i or '').strip().replace('-', '').upper()
        if i and i != 'NONE' and i.zfill(8) in need: res.setdefault(i.zfill(8), []).append(rec)
json.dump({'meta': {'source': 'Elsevier Scopus Source List', 'sheet': built, 'file': os.path.basename(src), 'issn': len(res)}, 'sources': res}, open(out, 'w', encoding='utf8'), ensure_ascii=False, separators=(',', ':'))
print(len(res), 'ISSN ->', out)
